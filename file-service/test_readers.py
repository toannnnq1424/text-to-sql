import datetime as dt
import io
import json
from decimal import Decimal
from pathlib import Path

import fastavro
import pyarrow as pa
import pyarrow.feather as feather
import pyarrow.orc as orc
import pyarrow.parquet as parquet
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import app
from readers import InvalidFile, canonical, parse


def parsed(tmp_path, extension, data, **options):
    source = tmp_path / ("input." + extension)
    source.write_bytes(data)
    spool = tmp_path / "rows.ndjson"
    metadata = parse(source, extension, spool, **options)
    rows = [json.loads(line)["row"] for line in spool.read_text().splitlines()]
    return metadata, rows


@pytest.mark.parametrize(
    "extension,delimiter", [("csv", ","), ("tsv", "\t"), ("txt", ";"), ("txt", "|")]
)
def test_delimited_full_inference_and_quoted_values(tmp_path, extension, delimiter):
    content = f'ID{delimiter}Tên nhạc{delimiter}Amount{delimiter}enabled\r\n001{delimiter}"a\nb"{delimiter}123.4500{delimiter}true\r\n002{delimiter}O\'Brien{delimiter}1.00{delimiter}false\r\n'
    metadata, rows = parsed(tmp_path, extension, content.encode())
    assert [c["dataType"] for c in metadata["columns"]] == [
        "STRING",
        "STRING",
        "DECIMAL",
        "BOOLEAN",
    ]
    assert metadata["columns"][1]["name"] == "ten_nhac"
    assert rows[0] == ["001", "a\nb", "123.4500", "true"]


def test_full_file_not_first_preview_inference(tmp_path):
    values = "value\n" + "1\n" * 25 + "last_is_text\n"
    metadata, rows = parsed(tmp_path, "csv", values.encode())
    assert metadata["columns"][0]["dataType"] == "STRING"
    assert len(rows) == 26


def test_utf16_exact_bigint_and_null(tmp_path):
    metadata, rows = parsed(tmp_path, "tsv", "id\tvalue\n9007199254740993\t\n".encode("utf-16"))
    assert rows[0][0] == "9007199254740993"
    assert metadata["columns"][0]["dataType"] == "BIGINT"
    assert metadata["columns"][1]["nullable"]


@pytest.mark.parametrize("extension", ["json", "jsonl", "ndjson"])
def test_json_union_nested_and_precision(tmp_path, extension):
    records = [
        {"id": 9007199254740993, "label": "Việt", "info": {"x": 1}},
        {"id": 2, "later": True},
    ]
    data = (
        json.dumps(records)
        if extension == "json"
        else "\n".join(json.dumps(row) for row in records)
    )
    metadata, rows = parsed(tmp_path, extension, data.encode())
    assert metadata["rowCount"] == 2
    assert rows[0][0] == "9007199254740993"
    assert json.loads(rows[0][2]) == {"x": 1}
    assert rows[0][3] is None


@pytest.mark.parametrize("extension", ["parquet", "feather", "arrow", "ipc", "orc"])
def test_native_formats_exact_decimal_and_temporal(tmp_path, extension):
    table = pa.table(
        {
            "id": pa.array([9007199254740993, 2], type=pa.int64()),
            "amount": pa.array([Decimal("123.4500"), None], type=pa.decimal128(20, 4)),
            "date": [dt.date(2026, 10, 20), dt.date(2026, 10, 21)],
        }
    )
    output = pa.BufferOutputStream()
    if extension == "parquet":
        parquet.write_table(table, output)
    elif extension == "feather":
        feather.write_feather(table, output)
    elif extension == "orc":
        orc.write_table(table, output)
    else:
        with pa.ipc.new_file(output, table.schema) as writer:
            writer.write_table(table)
    metadata, rows = parsed(tmp_path, extension, output.getvalue().to_pybytes())
    assert rows[0] == ["9007199254740993", "123.4500", "2026-10-20"]
    assert [c["dataType"] for c in metadata["columns"]] == ["BIGINT", "DECIMAL", "DATE"]


def test_avro_and_sheet_selection(tmp_path):
    buffer = io.BytesIO()
    fastavro.writer(
        buffer,
        {
            "type": "record",
            "name": "Data",
            "fields": [{"name": "id", "type": "long"}, {"name": "value", "type": "string"}],
        },
        [{"id": 1, "value": "nhạc"}],
    )
    assert parsed(tmp_path, "avro", buffer.getvalue())[1] == [["1", "nhạc"]]
    workbook = Workbook()
    workbook.active.append(["a"])
    workbook.active.append([1])
    sheet = workbook.create_sheet("Nhạc")
    sheet.append(["name", "quantity"])
    sheet.append(["Việt", 3])
    buffer = io.BytesIO()
    workbook.save(buffer)
    metadata, rows = parsed(tmp_path, "xlsx", buffer.getvalue(), sheet="Nhạc")
    assert metadata["sheets"] == ["Sheet", "Nhạc"]
    assert metadata["sheet"] == "Nhạc"
    assert rows == [["Việt", "3"]]


def test_formulas_duplicate_headers_and_bad_row_rejected(tmp_path):
    for data in [b"a,a\n1,2\n", b"a,b\n1,2,3\n", b"a\n", b"a\n\x00\n"]:
        with pytest.raises(InvalidFile):
            parsed(tmp_path, "csv", data)
    workbook = Workbook()
    workbook.active.append(["value"])
    workbook.active.append(["=1+2"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(InvalidFile, match="Formula"):
        parsed(tmp_path, "xlsx", buffer.getvalue())
    with pytest.raises(InvalidFile):
        parsed(tmp_path, "sql", b"DROP DATABASE whatever")


def test_auth_and_stream_complete_output(monkeypatch):
    monkeypatch.setenv("FILE_SERVICE_TOKEN", "x" * 40)
    client = TestClient(app)
    assert (
        client.post("/internal/v1/parse", files={"file": ("data.csv", b"a\n1\n")}).status_code
        == 401
    )
    response = client.post(
        "/internal/v1/parse",
        files={"file": ("data.csv", b"a\n1\n2\n")},
        headers={"Authorization": "Bearer " + "x" * 40},
    )
    assert response.status_code == 200
    lines = [json.loads(line) for line in response.text.splitlines()]
    assert lines[0]["rowCount"] == 2
    assert lines[1:] == [{"row": ["1"]}, {"row": ["2"]}]


def test_nested_json_numbers_and_limits():
    result = canonical({"amount": Decimal("123.4500"), "data": [Decimal("0.00000000001")]})
    assert '"amount":123.4500' in result
    assert json.loads(result, parse_float=Decimal)["amount"] == Decimal("123.4500")
    with pytest.raises(InvalidFile):
        canonical({"large": "a" * 17000})
    with pytest.raises(InvalidFile):
        canonical({"value": Decimal("NaN")})
