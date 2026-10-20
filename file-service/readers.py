"""Bounded local-file readers, adapted from Apache-2.0 object-storage dispatch.

See third-party/NOTICE.md. No reader accepts network URLs or user server paths.
"""

import csv
import datetime as dt
import io
import json
import math
import re
import unicodedata
import zipfile
from decimal import Decimal
from pathlib import Path

MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_CELL = 16_384
MAX_CANONICAL = 75 * 1024 * 1024
FORMATS = {"csv", "tsv", "txt", "json", "jsonl", "ndjson", "xlsx", "parquet", "arrow", "feather", "ipc", "orc", "avro"}


class InvalidFile(ValueError):
    pass


def encoded(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def canonical(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str, allow_nan=False)
    if isinstance(value, dt.datetime):
        if value.tzinfo:
            value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
        return value.isoformat(sep="T")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bytes):
        raise InvalidFile("Binary cells cannot be mapped to text automatically.")
    if isinstance(value, float) and not math.isfinite(value):
        raise InvalidFile("Non-finite numeric cells are unsupported.")
    text = str(value)
    if "\x00" in text or len(text) > MAX_CELL:
        raise InvalidFile("Cell exceeds limit or contains a null byte.")
    return text


def value_type(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or str(value).lower() in {"true", "false"}:
        return "BOOLEAN"
    text = str(value)
    if re.fullmatch(r"[+-]?(0|[1-9][0-9]*)", text):
        number = int(text)
        if -(2**63) <= number < 2**63:
            return "BIGINT"
    # Leading zeros are identifiers, not quantities.
    if re.fullmatch(r"[+-]?0[0-9]+(?:\.[0-9]+)?", text):
        return "STRING"
    if re.fullmatch(r"[+-]?(?:[0-9]+\.[0-9]*|[0-9]*\.[0-9]+|[0-9]+[eE][+-]?[0-9]+)", text):
        number = Decimal(text)
        _, digits, exponent = number.as_tuple()
        if max(len(digits) + exponent, 0) <= 28 and max(-exponent, 0) <= 10:
            return "DECIMAL"
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            dt.date.fromisoformat(text)
            return "DATE"
        if re.match(r"^\d{4}-\d{2}-\d{2}[T ]", text):
            parsed = dt.datetime.fromisoformat(text.replace("Z", "+00:00"))
            if 1000 <= parsed.year <= 9999:
                return "TIMESTAMP"
    except ValueError:
        pass
    return "STRING"


def merge_type(left, right):
    if left is None:
        return right
    if right is None or left == right:
        return left
    if {left, right} <= {"BIGINT", "DECIMAL"}:
        return "DECIMAL"
    if {left, right} <= {"DATE", "TIMESTAMP"}:
        return "TIMESTAMP"
    return "STRING"


def safe_names(headers):
    if not headers or len(headers) > MAX_COLUMNS:
        raise InvalidFile("File must contain 1–200 columns.")
    if any(not isinstance(h, str) or not h.strip() or len(h) > 255 for h in headers):
        raise InvalidFile("Column headers must be nonempty text, at most 255 characters.")
    if len(set(headers)) != len(headers):
        raise InvalidFile("Duplicate source headers must be corrected before upload.")
    names = []
    used = set()
    for header in headers:
        ascii_name = unicodedata.normalize("NFKD", header.replace("đ", "d").replace("Đ", "D"))
        base = re.sub(r"[^a-z0-9_]+", "_", ascii_name.encode("ascii", "ignore").decode().lower()).strip("_")[:44]
        if not base or base[0].isdigit():
            base = "col_" + base
        name = base
        suffix = 1
        while name.casefold() in used:
            suffix += 1
            name = f"{base}_{suffix}"
        names.append(name)
        used.add(name.casefold())
    return names


def _delimited(path, extension, delimiter):
    raw = path.read_bytes()
    encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    try:
        text = raw.decode(encoding)
    except UnicodeError as exc:
        raise InvalidFile("Use UTF-8 or UTF-16 with BOM for delimited text.") from exc
    if delimiter is not None and delimiter not in {",", "\t", ";", "|"}:
        raise InvalidFile("Unsupported delimiter.")
    if delimiter is None:
        if extension == "tsv":
            delimiter = "\t"
        else:
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",\t;|").delimiter
            except csv.Error:
                delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise InvalidFile("File is empty.") from exc
    return headers, reader, {"warnings": [f"Encoding: {encoding}; delimiter: {repr(delimiter)}"]}


def _records(records):
    rows = []
    headers = []
    for record in records:
        if not isinstance(record, dict):
            raise InvalidFile("JSON/Avro rows must be objects.")
        for key in record:
            if key not in headers:
                headers.append(key)
        if len(headers) > MAX_COLUMNS or len(rows) >= MAX_ROWS:
            raise InvalidFile("Row or column limit exceeded.")
        rows.append(record)
    return headers, ([row.get(h) for h in headers] for row in rows), {}


def _xlsx(path, sheet):
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 150 * 1024 * 1024:
            raise InvalidFile("Workbook expanded size exceeds limit.")
        for entry in entries:
            name = entry.filename.lower()
            if entry.file_size > 1024 * 1024 and entry.file_size > max(entry.compress_size, 1) * 250:
                raise InvalidFile("Workbook compression ratio exceeds limit.")
            if "vbaproject" in name or "externallinks/" in name:
                raise InvalidFile("Macros and external workbook references are not accepted.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    sheets = workbook.sheetnames
    selected = sheet or sheets[0]
    if selected not in sheets:
        workbook.close()
        raise InvalidFile("Unknown worksheet.")
    worksheet = workbook[selected]
    if worksheet.max_column > MAX_COLUMNS or worksheet.max_row > MAX_ROWS + 1:
        workbook.close()
        raise InvalidFile("Workbook row or column limit exceeded.")
    iterator = worksheet.iter_rows()
    headers = [cell.value for cell in next(iterator, [])]

    def rows():
        try:
            for row in iterator:
                if any(cell.data_type == "f" for cell in row):
                    raise InvalidFile("Formula cells must be replaced with values before upload.")
                if any(cell.value is not None for cell in row):
                    yield [cell.value for cell in row]
        finally:
            workbook.close()

    return headers, rows(), {"sheets": sheets, "sheet": selected}


def _native(path, extension):
    # Local-file dispatch adapted from the reviewed Apache source; reader calls
    # preserve Arrow values rather than converting exact decimals through floats.
    import pyarrow as pa
    import pyarrow.feather as feather
    import pyarrow.orc as orc
    import pyarrow.parquet as parquet

    if extension == "parquet":
        source = parquet.ParquetFile(path)
        if source.metadata.num_rows > MAX_ROWS:
            raise InvalidFile("Row limit exceeded.")
        batches = source.iter_batches(batch_size=1000)
        headers = source.schema_arrow.names
    elif extension == "orc":
        source = orc.ORCFile(path)
        if source.nrows > MAX_ROWS:
            raise InvalidFile("Row limit exceeded.")
        headers = source.schema.names
        batches = (source.read_stripe(i) for i in range(source.nstripes))
    else:
        if extension == "feather":
            table = feather.read_table(path, memory_map=True)
            headers = table.column_names
            batches = table.to_batches(max_chunksize=1000)
        else:
            mapped = pa.memory_map(str(path), "r")
            try:
                source = pa.ipc.open_file(mapped)
                headers = source.schema.names
                batches = (source.get_batch(i) for i in range(source.num_record_batches))
            except pa.ArrowInvalid:
                mapped.seek(0)
                source = pa.ipc.open_stream(mapped)
                headers = source.schema.names
                batches = iter(source)

    def rows():
        for batch in batches:
            for record in batch.to_pylist():
                yield [record[h] for h in headers]

    return headers, rows(), {}


def parse(path: Path, extension: str, spool: Path, sheet=None, delimiter=None):
    if extension not in FORMATS:
        raise InvalidFile("Unsupported file format; SQL files use SQL Import.")
    if extension in {"csv", "tsv", "txt"}:
        headers, rows, info = _delimited(path, extension, delimiter)
    elif extension in {"json", "ndjson", "jsonl"}:
        text = path.read_text(encoding="utf-8-sig")
        if extension == "json":
            records = json.loads(text, parse_float=Decimal)
            if not isinstance(records, list):
                raise InvalidFile("JSON must be an array of objects; use NDJSON for records.")
        else:
            records = (json.loads(line, parse_float=Decimal) for line in text.splitlines() if line.strip())
        headers, rows, info = _records(records)
    elif extension == "xlsx":
        headers, rows, info = _xlsx(path, sheet)
    elif extension == "avro":
        import fastavro

        with path.open("rb") as stream:
            headers, rows, info = _records(fastavro.reader(stream))
    else:
        headers, rows, info = _native(path, extension)
    names = safe_names(headers)
    types = [None] * len(headers)
    nulls = [False] * len(headers)
    count = 0
    size = 0
    with spool.open("w", encoding="utf-8") as output:
        for row in rows:
            count += 1
            if count > MAX_ROWS or len(row) != len(headers):
                raise InvalidFile(f"Invalid row width or row limit at row {count}.")
            values = []
            for i, value in enumerate(row):
                value = canonical(value)
                nulls[i] |= value is None or value == ""
                types[i] = merge_type(types[i], value_type(value))
                values.append(value)
            line = encoded({"row": values}) + "\n"
            size += len(line.encode("utf-8"))
            if size > MAX_CANONICAL:
                raise InvalidFile("Canonical file size exceeds limit.")
            output.write(line)
    if not count:
        raise InvalidFile("File contains no data rows.")
    return {
        "fileType": extension.upper(),
        "rowCount": count,
        "columns": [
            {"sourceName": header, "name": name, "dataType": kind or "STRING", "nullable": nullable}
            for header, name, kind, nullable in zip(headers, names, types, nulls)
        ],
        "sheets": info.get("sheets", []),
        "warnings": info.get("warnings", []),
        **({"sheet": info["sheet"]} if "sheet" in info else {}),
    }
