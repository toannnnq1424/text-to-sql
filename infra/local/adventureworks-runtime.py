#!/usr/bin/env python3
"""Prepare pinned, real-source files and independent oracles for recorded UI acceptance.

This runner performs no API calls, registry writes, database operations or browser actions.
Files stay under ignored .local/adventureworks/ui-files; the product UI owns ingestion.
Run with file-service/.venv/bin/python for the optional binary format exports.
"""

import argparse
import csv
import hashlib
import importlib.util
import json
import sys
import uuid
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
LOCAL = ROOT / ".local/adventureworks/ui-files"
sys.path.insert(0, str(ROOT / "data/scripts"))
SPEC = importlib.util.spec_from_file_location(
    "adventureworks_source", ROOT / "data/scripts/adventureworks.py"
)
SOURCE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(SOURCE)


def persist(path, document):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(document, indent=2, ensure_ascii=False) + "\n")
    path.chmod(0o600)


def decimal_sum(rows, field):
    return format(
        sum((Decimal(row[field]) for row in rows if row[field] is not None), Decimal(0)), "f"
    )


def groups(rows, key, amount):
    values = defaultdict(Decimal)
    for row in rows:
        values[row[key]] += Decimal(row[amount])
    return {str(key): format(value, "f") for key, value in sorted(values.items())}


def oracle(rows):
    products = rows["DimProduct"]
    facts = rows["FactInternetSales"]
    territories = {row["SalesTerritoryKey"]: row for row in rows["DimSalesTerritory"]}
    countries = defaultdict(Decimal)
    for row in facts:
        countries[territories[row["SalesTerritoryKey"]]["SalesTerritoryCountry"]] += Decimal(
            row["SalesAmount"]
        )
    return {
        "origin": "Independently aggregated pinned Microsoft source records, never UI/API output",
        "sourceSha256": SOURCE.SOURCE["sha256"],
        "headers": {table: list(records[0]) for table, records in rows.items()},
        "counts": {table: len(records) for table, records in rows.items()},
        "DimProduct": {
            "rows": len(products),
            "listPriceSum": decimal_sum(products, "ListPrice"),
            "standardCostSum": decimal_sum(products, "StandardCost"),
            "colorCounts": dict(Counter(row["Color"] for row in products)),
            "nullCounts": {
                field: sum(row[field] is None for row in products) for field in products[0]
            },
            "byProductKey": {row["ProductKey"]: row for row in products},
        },
        "FactInternetSales": {
            "rows": len(facts),
            "totalRevenue": decimal_sum(facts, "SalesAmount"),
            "totalQuantity": decimal_sum(facts, "OrderQuantity"),
            "totalProductCost": decimal_sum(facts, "TotalProductCost"),
            "byProductKey": groups(facts, "ProductKey", "SalesAmount"),
            "bySalesTerritoryKey": groups(facts, "SalesTerritoryKey", "SalesAmount"),
            "byOrderDateKey": groups(facts, "OrderDateKey", "SalesAmount"),
            "byCountry": {country: format(amount, "f") for country, amount in countries.items()},
        },
    }


def export_csv(path, rows, delimiter=","):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)
    with path.open(encoding="utf-8", newline="") as stream:
        reread = list(csv.DictReader(stream, delimiter=delimiter))
    assert len(reread) == len(rows)
    assert all(
        restored == {key: "" if value is None else value for key, value in original.items()}
        for restored, original in zip(reread, rows)
    )


def binary_exports(rows):
    import fastavro
    import pyarrow as pa
    import pyarrow.feather as feather
    import pyarrow.orc as orc
    import pyarrow.parquet as parquet
    from openpyxl import Workbook, load_workbook

    paths = []
    table = pa.Table.from_pylist(rows)
    for extension in ("parquet", "feather", "arrow", "ipc", "orc"):
        path = LOCAL / ("DimProduct." + extension)
        if extension == "parquet":
            parquet.write_table(table, path)
            restored = parquet.read_table(path)
        elif extension == "feather":
            feather.write_feather(table, path)
            restored = feather.read_table(path)
        elif extension == "orc":
            orc.write_table(table, path)
            restored = orc.read_table(path)
        else:
            with (
                pa.OSFile(str(path), "wb") as output,
                pa.ipc.new_file(output, table.schema) as writer,
            ):
                writer.write_table(table)
            with pa.memory_map(str(path), "r") as source:
                restored = pa.ipc.open_file(source).read_all()
        assert restored.to_pylist() == rows
        paths.append(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(list(rows[0]))
    for row in rows:
        sheet.append(list(row.values()))
    path = LOCAL / "DimProduct.xlsx"
    workbook.save(path)
    restored = load_workbook(path, read_only=True, data_only=True)
    try:
        contents = list(restored.active.values)
        assert list(contents[0]) == list(rows[0])
        assert [dict(zip(contents[0], row)) for row in contents[1:]] == rows
    finally:
        restored.close()
    paths.append(path)
    path = LOCAL / "DimProduct.avro"
    schema = {
        "type": "record",
        "name": "AdventureWorksProductProjection",
        "fields": [{"name": name, "type": ["null", "string"]} for name in rows[0]],
    }
    with path.open("wb") as output:
        fastavro.writer(output, schema, rows)
    with path.open("rb") as stream:
        assert list(fastavro.reader(stream)) == rows
    paths.append(path)
    return paths


def prepare(binary=False):
    rows = SOURCE.source_data()
    LOCAL.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, records in rows.items():
        path = LOCAL / (name + ".csv")
        export_csv(path, records)
        paths.append((path, name, len(records)))
    product = rows["DimProduct"]
    for extension, delimiter in (("tsv", "\t"), ("txt", ";")):
        path = LOCAL / ("DimProduct." + extension)
        export_csv(path, product, delimiter)
        paths.append((path, "DimProduct", len(product)))
    path = LOCAL / "DimProduct.json"
    persist(path, product)
    assert json.loads(path.read_text()) == product
    paths.append((path, "DimProduct", len(product)))
    for extension in ("jsonl", "ndjson"):
        path = LOCAL / ("DimProduct." + extension)
        path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in product))
        assert [json.loads(line) for line in path.read_text().splitlines()] == product
        paths.append((path, "DimProduct", len(product)))
    if binary:
        paths.extend((path, "DimProduct", len(product)) for path in binary_exports(product))
    manifest = {
        "sourceUrl": SOURCE.SOURCE["url"],
        "sourceSha256": SOURCE.SOURCE["sha256"],
        "license": SOURCE.SOURCE["license"],
        "projection": SOURCE.SOURCE["selection"],
        "apiMutations": False,
        "browserActions": False,
        "exports": [],
    }
    for path, name, count in paths:
        size = path.stat().st_size
        if size > 25 * 1024 * 1024 or count > 100000:
            raise ValueError("Prepared file exceeds the real application ingestion limits.")
        manifest["exports"].append(
            {
                "path": str(path),
                "table": name,
                "rows": count,
                "columns": len(rows[name][0]),
                "bytes": size,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "roundTrip": "PASS",
            }
        )
    persist(LOCAL / "manifest.json", manifest)
    persist(LOCAL / "oracle.json", oracle(rows))
    print(json.dumps({"status": "PREPARED", "files": len(paths), "directory": str(LOCAL)}))


def prepare_sql():
    products = SOURCE.source_data()["DimProduct"]
    suffix = uuid.uuid4().hex[:10]
    fields = (
        ("productkey", "ProductKey", "INTEGER"),
        ("productname", "EnglishProductName", "STRING"),
        ("color", "Color", "STRING"),
        ("listprice", "ListPrice", "DECIMAL"),
    )
    manifest = {
        "sourceUrl": SOURCE.SOURCE["url"],
        "sourceSha256": SOURCE.SOURCE["sha256"],
        "license": SOURCE.SOURCE["license"],
        "rows": len(products),
        "projection": [
            {"name": name, "sourceColumn": source, "type": kind} for name, source, kind in fields
        ],
        "nullCounts": {
            name: sum(row[source] is None for row in products) for name, source, _ in fields
        },
        "listPriceSum": decimal_sum(products, "ListPrice"),
        "byProductKey": {
            row["ProductKey"]: {name: row[source] for name, source, _ in fields} for row in products
        },
        "apiMutations": False,
        "exports": [],
    }
    LOCAL.mkdir(parents=True, exist_ok=True)
    for engine, label in (("POSTGRESQL", "pg"), ("MYSQL", "mysql"), ("SQL_SERVER", "mssql")):
        name = "readiness_sql_product_" + label + "_" + suffix
        table = SOURCE.quote(engine, "nla_import") + "." + SOURCE.quote(engine, name)
        quoted = [SOURCE.quote(engine, field) for field, _, _ in fields]
        string_type = "NVARCHAR(255)" if engine == "SQL_SERVER" else "VARCHAR(255)"
        definitions = (
            quoted[0] + " INTEGER PRIMARY KEY",
            quoted[1] + " " + string_type + " NOT NULL",
            quoted[2] + " " + string_type,
            quoted[3] + " DECIMAL(19,4)",
        )
        statements = ["CREATE TABLE " + table + " (\n  " + ",\n  ".join(definitions) + "\n);"]
        values = [
            "("
            + ", ".join(
                SOURCE.literal(engine, row[source], {"type": kind}) for _, source, kind in fields
            )
            + ")"
            for row in products
        ]
        statements.append(
            "INSERT INTO "
            + table
            + " ("
            + ", ".join(quoted)
            + ") VALUES\n  "
            + ",\n  ".join(values)
            + ";"
        )
        path = LOCAL / ("DimProduct-" + label + ".sql")
        path.write_text("\n".join(statements) + "\n", encoding="utf-8")
        if path.stat().st_size > 2 * 1024 * 1024 or len(products) > 10000:
            raise ValueError("SQL projection exceeds application import bounds.")
        manifest["exports"].append(
            {
                "engine": engine,
                "table": name,
                "schema": "nla_import",
                "path": str(path),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
        )
    persist(LOCAL / "sql-import-manifest.json", manifest)
    print(
        json.dumps(
            {
                "status": "PREPARED",
                "sqlFiles": len(manifest["exports"]),
                "manifest": str(LOCAL / "sql-import-manifest.json"),
            }
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("prepare-files", "prepare-sql"))
    parser.add_argument("--binary", action="store_true")
    arguments = parser.parse_args()
    if arguments.command == "prepare-files":
        prepare(binary=arguments.binary)
    else:
        prepare_sql()
